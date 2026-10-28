"""Тести прийому даних з WMS.

Скрипт на машині біля WMS сам шле ruchy на /api/wms-import. Логіну він не
має, тож замість сесії несе спільний ключ у заголовку. Екран звіту бере
найсвіжішу пачку сам.
"""
import inspect
import json
from datetime import datetime, timedelta, timezone

import pytest
from fastapi import HTTPException

from app.dependencies.admin import require_admin
from models.db_wms_import import WmsImportDB
from routers.admin.report_wms import wms_report_latest
from routers.wms_import import wms_import, zapisz_paczke

pytestmark = pytest.mark.asyncio

KLUCZ = "tajny-klucz-testowy"


class _Request:
    """Tyle z żądania, ile czyta endpoint."""

    def __init__(self, body: bytes):
        self._body = body

    async def body(self) -> bytes:
        return self._body


def _paczka(ile: int = 2) -> bytes:
    teraz = datetime.now()
    wiersze = [
        {"login": f"P-TEST{i}", "code": "Pick",
         "dstamp": (teraz - timedelta(minutes=i)).strftime("%y/%m/%d %H:%M:%S,000000000")}
        for i in range(ile)
    ]
    return json.dumps({"items": wiersze}).encode()


@pytest.fixture
def klucz(monkeypatch):
    monkeypatch.setenv("WMS_IMPORT_KEY", KLUCZ)
    return KLUCZ


# ---------------------------------------------------------------------------
# Klucz zamiast logowania
# ---------------------------------------------------------------------------
async def test_bez_klucza_nie_wejdzie(db_session, klucz):
    with pytest.raises(HTTPException) as blad:
        await wms_import(request=_Request(_paczka()), x_pinokio_key=None,
                         x_pinokio_origin=None, db=db_session)

    assert blad.value.status_code == 401


async def test_zly_klucz_nie_wejdzie(db_session, klucz):
    with pytest.raises(HTTPException) as blad:
        await wms_import(request=_Request(_paczka()), x_pinokio_key="nie-ten",
                         x_pinokio_origin=None, db=db_session)

    assert blad.value.status_code == 401


async def test_bez_klucza_na_serwerze_endpoint_milczy(db_session, monkeypatch):
    """Lepiej, żeby import zamilkł, niż żeby każdy mógł tu wrzucić cokolwiek."""
    monkeypatch.delenv("WMS_IMPORT_KEY", raising=False)

    with pytest.raises(HTTPException) as blad:
        await wms_import(request=_Request(_paczka()), x_pinokio_key="cokolwiek",
                         x_pinokio_origin=None, db=db_session)

    assert blad.value.status_code == 503


async def test_endpoint_nie_wymaga_logowania():
    """Skrypt chodzi bez człowieka — nie ma czym się zalogować."""
    parametry = inspect.signature(wms_import).parameters

    assert "current_user" not in parametry
    assert not any("require_" in str(p.default) for p in parametry.values())


# ---------------------------------------------------------------------------
# Zapis paczki
# ---------------------------------------------------------------------------
async def test_paczka_laduje_w_bazie(db_session, klucz):
    odpowiedz = await wms_import(request=_Request(_paczka(3)), x_pinokio_key=KLUCZ,
                                 x_pinokio_origin="WMS-SRV", db=db_session)

    assert odpowiedz["status"] == "ok"
    assert odpowiedz["rows"] == 3

    from sqlalchemy import select
    paczka = (await db_session.execute(select(WmsImportDB))).scalars().one()
    assert paczka.rows_total == 3
    assert paczka.source == "push"
    assert paczka.origin == "WMS-SRV"
    assert len(paczka.payload) == 3


async def test_pusta_paczka_jest_odrzucona(db_session, klucz):
    with pytest.raises(HTTPException) as blad:
        await wms_import(request=_Request(b""), x_pinokio_key=KLUCZ,
                         x_pinokio_origin=None, db=db_session)

    assert blad.value.status_code == 400


async def test_smieci_zamiast_json_nie_przechodza(db_session, klucz):
    with pytest.raises(HTTPException) as blad:
        await wms_import(request=_Request(b"to nie jest json"), x_pinokio_key=KLUCZ,
                         x_pinokio_origin=None, db=db_session)

    assert blad.value.status_code == 400


async def test_zapisujemy_surowe_ruchy_a_nie_wynik(db_session):
    """Wydania sprzętu dochodzą po imporcie, więc zestawienie liczy się przy
    otwieraniu raportu — w bazie zostają same ruchy."""
    paczka = await zapisz_paczke(db_session, _paczka(2), source="push")

    assert {"login", "code", "dstamp"} == set(paczka.payload[0])
    assert paczka.period_from <= paczka.period_to


# ---------------------------------------------------------------------------
# Ekran bierze najnowszą paczkę
# ---------------------------------------------------------------------------
async def test_ekran_pokazuje_ostatni_import(db_session):
    await zapisz_paczke(db_session, _paczka(2), source="push", origin="STARY")
    await zapisz_paczke(db_session, _paczka(4), source="push", origin="NOWY")

    raport = await wms_report_latest(db=db_session, user=None)

    assert raport["totals"]["rows_total"] == 4
    assert raport["import"]["origin"] == "NOWY"
    assert raport["import"]["source"] == "push"


async def test_brak_importu_to_nie_awaria(db_session):
    with pytest.raises(HTTPException) as blad:
        await wms_report_latest(db=db_session, user=None)

    assert blad.value.status_code == 404


async def test_ostatni_import_tylko_dla_admina():
    guard = inspect.signature(wms_report_latest).parameters["user"].default.dependency

    assert guard is require_admin


async def test_czas_z_paczki_nie_gubi_strefy(db_session):
    """Znacznik zapisujemy w ISO ze strefą, żeby drugi raz jej nie zgadywać."""
    paczka = await zapisz_paczke(db_session, _paczka(1), source="push")

    zapisany = datetime.fromisoformat(paczka.payload[0]["dstamp"])
    assert zapisany.tzinfo is not None
    assert zapisany.utcoffset() in (timedelta(hours=1), timedelta(hours=2))


async def test_plik_xlsx_z_ostatniego_importu(db_session):
    """Przy danych ze skryptu nie ma czego wgrywać, a plik ma się pobrać."""
    from io import BytesIO

    from openpyxl import load_workbook

    from routers.admin.report_wms import wms_report_latest_xlsx

    await zapisz_paczke(db_session, _paczka(3), source="push", origin="WMS")

    odpowiedz = await wms_report_latest_xlsx(db=db_session, user=None)
    arkusze = load_workbook(BytesIO(odpowiedz.body)).sheetnames

    assert arkusze == ["Bez sprzętu", "Nieznane loginy", "Zakres"]
    assert "attachment" in odpowiedz.headers["Content-Disposition"]


async def test_plik_z_importu_tylko_dla_admina():
    from routers.admin.report_wms import wms_report_latest_xlsx

    guard = inspect.signature(wms_report_latest_xlsx).parameters["user"].default.dependency

    assert guard is require_admin
