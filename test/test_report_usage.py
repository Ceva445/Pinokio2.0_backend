"""Тести звіту «Wykorzystanie sprzętu».

Звіт відповідає на два питання одним набором: скільки пристрій уже в руках у
працівника (від останньої реєстрації) і скільки лежить без діла (від останнього
повернення). Час рахує сервер — вкладка може висіти годинами, і її годинник тут
ні до чого.
"""
import inspect
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from io import BytesIO

from app.dependencies.admin import require_admin
from models.db_device import DeviceDB, DeviceType
from models.db_device_status import DeviceStatusDB
from models.db_employee import EmployeeDB
from models.db_site import SiteDB
from models.db_transaction import TransactionDB, TransactionType
from routers.admin.report_usage import (
    build_workbook,
    content_disposition,
    format_czas,
    report_file_name,
    usage_report,
    usage_report_xlsx,
    zbierz,
)

pytestmark = pytest.mark.asyncio


def _kiedys(godzin):
    return datetime.now(tz=timezone.utc) - timedelta(hours=godzin)


@pytest.fixture
async def magazyn(db_session):
    site = SiteDB(name="STOCK")
    status = DeviceStatusDB(name="WORK")
    anna = EmployeeDB(last_name="Nowak", first_name="Anna", rfid="r-anna",
                      company="Demo", wms_login="A-NOWAK")
    db_session.add_all([site, status, anna])
    await db_session.commit()

    anna.site_id = site.id

    # W rękach od pięciu godzin.
    w_uzyciu = DeviceDB(name="TERM003", rfid="r-t3", serial_number="sn-t3",
                        type=DeviceType.scanner, employee_id=anna.id,
                        site_id=site.id, status_id=status.id)
    # Zwrócony dobę temu i od tamtej pory leży.
    wolne = DeviceDB(name="TERM004", rfid="r-t4", serial_number="sn-t4",
                     type=DeviceType.printer, site_id=site.id)
    # Nigdy nie wydany — to nie jest "zero minut".
    nowy = DeviceDB(name="TERM005", rfid="r-t5", serial_number="sn-t5",
                    type=DeviceType.scanner)
    db_session.add_all([w_uzyciu, wolne, nowy])
    await db_session.commit()

    db_session.add_all([
        # Stara rejestracja tego samego sprzętu — liczy się najnowsza.
        TransactionDB(timestamp=_kiedys(90), type=TransactionType.registered,
                      device_id=w_uzyciu.id, employee_id=anna.id),
        TransactionDB(timestamp=_kiedys(5), type=TransactionType.registered,
                      device_id=w_uzyciu.id, employee_id=anna.id),
        TransactionDB(timestamp=_kiedys(30), type=TransactionType.registered,
                      device_id=wolne.id, employee_id=anna.id),
        TransactionDB(timestamp=_kiedys(24), type=TransactionType.unregistered,
                      device_id=wolne.id),
    ])
    await db_session.commit()

    return {"anna": anna, "w_uzyciu": w_uzyciu, "wolne": wolne, "nowy": nowy}


async def _po_nazwie(db):
    return {d["name"]: d for d in await zbierz(db)}


# ---------------------------------------------------------------------------
# Skąd liczymy czas
# ---------------------------------------------------------------------------
async def test_sprzet_w_rekach_liczy_od_rejestracji(db_session, magazyn):
    d = (await _po_nazwie(db_session))["TERM003"]

    assert d["in_use"] is True
    assert d["employee"]["wms_login"] == "A-NOWAK"
    assert 4 * 60 + 55 < d["minutes"] < 5 * 60 + 5


async def test_liczy_sie_najnowsza_rejestracja(db_session, magazyn):
    """Ten sam sprzęt bywa wydawany wielokrotnie — 90 godzin temu nas nie
    interesuje, skoro ostatnia rejestracja była pięć godzin temu."""
    d = (await _po_nazwie(db_session))["TERM003"]

    assert d["minutes"] < 6 * 60


async def test_wolny_sprzet_liczy_od_zwrotu(db_session, magazyn):
    d = (await _po_nazwie(db_session))["TERM004"]

    assert d["in_use"] is False
    assert d["employee"] is None
    assert 23 * 60 + 55 < d["minutes"] < 24 * 60 + 5


async def test_sprzet_bez_historii_nie_ma_zera(db_session, magazyn):
    """Nigdy nie wydany sprzęt nie stoi "0 minut" — on po prostu nie ma czasu."""
    d = (await _po_nazwie(db_session))["TERM005"]

    assert d["minutes"] is None
    assert d["since"] is None


# ---------------------------------------------------------------------------
# Format czasu
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("minuty, oczekiwane", [
    (None, "—"),
    (0, "0min"),
    (45, "45min"),
    (60, "1h 00min"),
    (614, "10h 14min"),
    (1876, "31h 16min"),
    (4320, "3d 0h"),        # powyżej trzech dób same godziny nic nie mówią
    (88086, "61d 4h"),
])
async def test_czas_po_ludzku(minuty, oczekiwane):
    assert format_czas(minuty) == oczekiwane


# ---------------------------------------------------------------------------
# Plik
# ---------------------------------------------------------------------------
async def test_plik_ma_nazwe_raportu_i_date():
    from datetime import date

    assert report_file_name(date(2026, 9, 11)) == "Wykorzystanie_sprzętu_2026-09-11.xlsx"


async def test_naglowek_pliku_przezyje_polskie_znaki():
    """Nagłówki HTTP jadą jako latin-1 — samo filename= z "sprzętu" wywracało
    pobieranie na 500."""
    naglowek = content_disposition("Wykorzystanie_sprzętu_2026-09-11.xlsx")

    naglowek.encode("latin-1")           # ma się zakodować bez wyjątku
    assert 'filename="Wykorzystanie_sprzetu_2026-09-11.xlsx"' in naglowek
    assert "filename*=UTF-8''" in naglowek


async def test_plik_ma_te_same_wiersze_co_ekran(db_session, magazyn):
    devices = await zbierz(db_session)
    sheet = load_workbook(BytesIO(build_workbook(devices))).active

    assert sheet.max_row == len(devices) + 1          # + nagłówek
    assert sheet["A1"].value == "Urządzenie"
    assert sheet["I1"].value == "Czas"


async def test_w_pliku_najdluzsze_na_gorze(db_session, magazyn):
    sheet = load_workbook(BytesIO(build_workbook(await zbierz(db_session)))).active
    nazwy = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]

    # Doba postoju przed pięcioma godzinami użycia, a sprzęt bez historii na końcu.
    assert nazwy == ["TERM004", "TERM003", "TERM005"]


async def test_sprzet_bez_historii_nie_udaje_zera_w_pliku(db_session, magazyn):
    sheet = load_workbook(BytesIO(build_workbook(await zbierz(db_session)))).active
    ostatni = sheet.max_row

    assert sheet.cell(row=ostatni, column=1).value == "TERM005"
    assert sheet.cell(row=ostatni, column=9).value == "—"
    assert sheet.cell(row=ostatni, column=10).value in (None, "")


# ---------------------------------------------------------------------------
# Dostęp
# ---------------------------------------------------------------------------
async def test_raport_tylko_dla_admina():
    for endpoint in (usage_report, usage_report_xlsx):
        guard = inspect.signature(endpoint).parameters["user"].default.dependency
        assert guard is require_admin, endpoint.__name__


async def test_odpowiedz_niesie_stan_na_kiedy(db_session, magazyn):
    dane = await usage_report(db=db_session, user=None)

    assert dane["report_name"] == "Wykorzystanie sprzętu"
    assert len(dane["generated_at"]) == 16          # YYYY-MM-DD HH:MM
    assert len(dane["devices"]) == 3
