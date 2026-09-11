"""Raport: wykorzystanie sprzętu.

Odpowiada na dwa pytania naraz. Ile sprzęt jest u pracownika — licząc od chwili,
gdy go zarejestrował. I ile leży bez pracy — licząc od ostatniego zwrotu.

Czas liczy serwer, nie przeglądarka: karta potrafi stać otwarta godzinami, a
zegar w niej bywa własny. Minuty jadą gotowe, żeby ekran i plik pokazały to samo.
"""
from datetime import date, datetime, timezone
from io import BytesIO
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.dependencies.admin import require_admin
from db.session import get_db
from models.db_device import DeviceDB, DeviceType
from models.db_employee import EmployeeDB
from models.db_transaction import TransactionDB, TransactionType

router = APIRouter(
    prefix="/admin/api/reports",
    tags=["Admin Reports"]
)

REPORT_NAME = "Wykorzystanie sprzętu"
REPORT_TZ = ZoneInfo("Europe/Warsaw")

COLUMNS = ("Urządzenie", "Typ", "Site", "Status", "Stan", "Pracownik",
           "Osoba", "Od", "Czas", "Minuty")
COLUMN_WIDTHS = (18, 12, 12, 14, 12, 20, 26, 20, 14, 10)

TYP_PL = {"scanner": "skaner", "printer": "drukarka"}


def format_czas(minuty: int | None) -> str:
    """Godziny i minuty — tak magazyn mówi o czasie. Powyżej trzech dób same
    godziny przestają cokolwiek znaczyć, więc wtedy dochodzą dni."""
    if minuty is None:
        return "—"
    godziny, minuta = divmod(int(minuty), 60)
    if godziny >= 72:
        return f"{godziny // 24}d {godziny % 24}h"
    return f"{godziny}h {minuta:02d}min" if godziny else f"{minuta}min"


def _aware(moment: datetime | None) -> datetime | None:
    """Znacznik ze strefą. Postgres oddaje je zonowane, ale nie każdy silnik —
    sqlite w testach zwraca gołe. Bez tego odejmowanie wywala się na 500."""
    if moment is None:
        return None
    return moment if moment.tzinfo else moment.replace(tzinfo=timezone.utc)


def _local(moment: datetime | None) -> datetime | None:
    moment = _aware(moment)
    if moment is None:
        return None
    return moment.astimezone(REPORT_TZ).replace(tzinfo=None)


async def zbierz(db: AsyncSession) -> list[dict]:
    """Każde urządzenie z czasem trwania jego obecnego stanu."""
    # Ostatnia rejestracja i ostatni zwrot per urządzenie — jednym przebiegiem,
    # zamiast pytać bazę osobno dla każdego z dwustu kilkudziesięciu sprzętów.
    ostatnie = (
        select(
            TransactionDB.device_id,
            func.max(TransactionDB.timestamp).filter(
                TransactionDB.type == TransactionType.registered
            ).label("rejestracja"),
            func.max(TransactionDB.timestamp).filter(
                TransactionDB.type == TransactionType.unregistered
            ).label("zwrot"),
        )
        .group_by(TransactionDB.device_id)
        .subquery()
    )

    stmt = (
        select(DeviceDB, ostatnie.c.rejestracja, ostatnie.c.zwrot)
        .outerjoin(ostatnie, ostatnie.c.device_id == DeviceDB.id)
        .options(
            selectinload(DeviceDB.employee).selectinload(EmployeeDB.site),
            selectinload(DeviceDB.site),
            selectinload(DeviceDB.status),
        )
        .order_by(DeviceDB.name)
    )

    teraz = datetime.now(tz=REPORT_TZ)
    wynik = []

    for device, rejestracja, zwrot in (await db.execute(stmt)).all():
        w_uzyciu = device.employee_id is not None
        od = _aware(rejestracja if w_uzyciu else zwrot)

        wynik.append({
            "name": device.name,
            "type": device.type.value,
            "site": device.site.name if device.site else None,
            "status": device.status.name if device.status else None,
            "enabled": device.enabled,
            "in_use": w_uzyciu,
            "employee": {
                "wms_login": device.employee.wms_login,
                "first_name": device.employee.first_name,
                "last_name": device.employee.last_name,
                "site": device.employee.site.name if device.employee.site else None,
            } if device.employee else None,
            "since": _local(od).strftime("%Y-%m-%d %H:%M") if od else None,
            # Brak historii to nie zero minut — takiego sprzętu nigdy nie wydano.
            "minutes": round((teraz - od).total_seconds() / 60) if od else None,
        })

    return wynik


@router.get("/usage")
async def usage_report(
    db: AsyncSession = Depends(get_db),
    user=Depends(require_admin),
):
    """Dane pod mini-dashboard. Filtrowanie zostaje po stronie ekranu —
    zbiór jest mały, a każdy klik w filtr nie musi wracać na serwer."""
    return {
        "report_name": REPORT_NAME,
        "generated_at": datetime.now(tz=REPORT_TZ).strftime("%Y-%m-%d %H:%M"),
        "devices": await zbierz(db),
    }


def _wiersz(d: dict) -> tuple:
    osoba = d["employee"]
    return (
        d["name"],
        TYP_PL.get(d["type"], d["type"]),
        d["site"] or "—",
        d["status"] or "—",
        "w użyciu" if d["in_use"] else "wolne",
        (osoba or {}).get("wms_login") or "—",
        " ".join(filter(None, [(osoba or {}).get("first_name"),
                               (osoba or {}).get("last_name")])) or "—",
        d["since"] or "—",
        format_czas(d["minutes"]),
        d["minutes"] if d["minutes"] is not None else "",
    )


def report_file_name(generated_on: date) -> str:
    return f"{REPORT_NAME.replace(' ', '_')}_{generated_on.isoformat()}.xlsx"


def content_disposition(file_name: str) -> str:
    """Nagłówek z nazwą pliku, która ma polskie znaki.

    Nagłówki HTTP jadą jako latin-1, a "sprzętu" się w tym nie mieści — samo
    filename= wywracało odpowiedź na 500. Właściwa nazwa idzie więc w filename*
    (RFC 5987), a obok zostaje wersja bez ogonków dla starszych przeglądarek.
    """
    ascii_name = (file_name.replace("ę", "e").replace("ó", "o").replace("ą", "a")
                  .replace("ś", "s").replace("ł", "l").replace("ż", "z")
                  .replace("ź", "z").replace("ć", "c").replace("ń", "n"))
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(file_name)}"


def build_workbook(devices: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wykorzystanie"

    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")

    # Najdłużej trwające na górze — plik ma zaczynać się od tego, co pilne.
    for d in sorted(devices, key=lambda x: x["minutes"] or -1, reverse=True):
        sheet.append(list(_wiersz(d)))

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@router.get("/usage.xlsx")
async def usage_report_xlsx(
    db: AsyncSession = Depends(get_db),
    user=Depends(require_admin),
):
    devices = await zbierz(db)
    file_name = report_file_name(datetime.now(REPORT_TZ).date())

    return Response(
        content=build_workbook(devices),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(file_name)},
    )
