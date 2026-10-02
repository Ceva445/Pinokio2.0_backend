"""Raport: ruchy w WMS kontra rejestracje sprzętu.

WMS wie, kto pracował — każda operacja (Pick, Receipt, Marshal…) zostaje tam
zapisana z loginem i czasem. Pinokio wie, kto pobrał skaner. Ten raport zbiera
jedno z drugim i pokazuje ludzi, którzy pracowali, a sprzętu u nas nie wzięli.

Za pracującego "u nas" uznajemy tego, kto w chwili swojego ruchu miał
wydany sprzęt — także wzięty przed początkiem pliku, bo zmiana nocna pobiera
skanery wieczorem.

Raport jest tylko dla administratora: zestawia pracę całego magazynu, nie
jednego działu.

Plik z WMS wgrywa człowiek — eksport w formacie JSON:

    [{"login": "P-PARPIEVN", "CODE": "Pick",
      "DSTAMP": "2026-10-01 04:00:58.000000 UTC"}, …]

Okres bierzemy z samego pliku: od najstarszego do najświeższego ruchu. Nikt
nie musi go wpisywać ani pamiętać, kiedy zrobiono eksport — plik sam mówi,
jaką zmianę opisuje.
"""
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.dependencies.admin import require_admin
from db.session import get_db
from models.db_employee import EmployeeDB
from models.db_transaction import TransactionDB, TransactionType
from routers.admin.report_usage import content_disposition

router = APIRouter(
    prefix="/admin/api/reports",
    tags=["Admin Reports"]
)

REPORT_NAME = "Praca w WMS bez pobranego sprzętu"
REPORT_TZ = ZoneInfo("Europe/Warsaw")

MAX_FILE_BYTES = 25 * 1024 * 1024


@dataclass(frozen=True)
class Ruch:
    """Jedna operacja z WMS."""
    login: str
    code: str
    moment: datetime         # zawsze ze strefą, w UTC


# ---------------------------------------------------------------------------
# Wejście
# ---------------------------------------------------------------------------
# "26/10/02 12:49:55,778394000" — rok dwucyfrowy, przecinek zamiast kropki
# i nanosekundy, których Python nie przyjmuje.
_KROTKA_DATA = re.compile(
    r"^(\d{2})/(\d{2})/(\d{2})[ T](\d{1,2}):(\d{2}):(\d{2})(?:[.,](\d+))?$"
)


def parse_dstamp(value: str) -> datetime:
    """Znacznik z WMS → datetime ze strefą.

    Eksport przychodzi w dwóch postaciach i każda ma inną strefę:

    * "2026-10-01 04:00:58.000000 UTC" — strefa napisana wprost;
    * "26/10/02 12:49:55,778394000" — czas lokalny magazynu, bez strefy.

    Tę drugą sprawdziliśmy danymi: przy czytaniu jej jako czasu lokalnego 142
    ze 149 ruchów wypada w chwili, gdy człowiek miał wydany sprzęt; przy
    czytaniu jako UTC — tylko 72. Dwugodzinne przesunięcie rozjeżdżało cały
    raport: ludzie wyglądali na pracujących po oddaniu skanera.
    """
    tekst = (value or "").strip()
    if not tekst:
        raise ValueError("pusty DSTAMP")

    dopasowanie = _KROTKA_DATA.match(tekst)
    if dopasowanie:
        rok, miesiac, dzien, godzina, minuta, sekunda, ulamek = dopasowanie.groups()
        mikrosekundy = int((ulamek or "0")[:6].ljust(6, "0"))
        return datetime(2000 + int(rok), int(miesiac), int(dzien),
                        int(godzina), int(minuta), int(sekunda), mikrosekundy,
                        tzinfo=REPORT_TZ)

    if tekst.upper().endswith("UTC"):
        tekst = tekst[:-3].strip()
    moment = datetime.fromisoformat(tekst.replace("T", " "))
    return moment if moment.tzinfo else moment.replace(tzinfo=timezone.utc)


def _wytnij_liste(tekst: str):
    """Lista z niepełnego pliku: od pierwszego [ do ostatniego ]."""
    poczatek, koniec = tekst.find("["), tekst.rfind("]")
    if poczatek == -1 or koniec < poczatek:
        raise HTTPException(400, "To nie jest plik z ruchami WMS (brak listy)")
    try:
        return json.loads(tekst[poczatek:koniec + 1])
    except json.JSONDecodeError as blad:
        raise HTTPException(400, f"To nie jest poprawny plik JSON: {blad}")


def parse_wms(raw: bytes) -> list[Ruch]:
    """Lista ruchów z wgranego pliku.

    Zły plik ma się skończyć czytelnym komunikatem, a nie pięćsetką: człowiek
    wgrywa to ręcznie i pomyłka jest normalną koleją rzeczy.
    """
    try:
        tekst = raw.decode("utf-8-sig")
    except UnicodeDecodeError as blad:
        raise HTTPException(400, f"Nie umiem odczytać pliku: {blad}")

    try:
        dane = json.loads(tekst)
    except json.JSONDecodeError:
        # Eksport z WMS bywa urwany: zaczyna się od ,"items": i kończy ]} —
        # jako całość to nie jest JSON. Bierzemy z niego samą listę.
        dane = _wytnij_liste(tekst)

    # Eksport bywa też opakowany: {"items": [...]} albo {"rows": [...]}.
    if isinstance(dane, dict):
        dane = next((v for v in dane.values() if isinstance(v, list)), None)
    if not isinstance(dane, list):
        raise HTTPException(400, "Plik ma zawierać listę ruchów z WMS")

    ruchy, bledne = [], 0
    for wiersz in dane:
        if not isinstance(wiersz, dict):
            bledne += 1
            continue
        login = str(wiersz.get("login") or wiersz.get("LOGIN") or "").strip()
        code = str(wiersz.get("CODE") or wiersz.get("code") or "").strip() or "—"
        try:
            moment = parse_dstamp(wiersz.get("DSTAMP") or wiersz.get("dstamp") or "")
        except ValueError:
            bledne += 1
            continue
        if not login:
            bledne += 1
            continue
        ruchy.append(Ruch(login=login, code=code, moment=moment))

    if not ruchy:
        raise HTTPException(
            400,
            "W pliku nie ma ani jednego czytelnego ruchu "
            "(potrzebne pola: login, CODE, DSTAMP)",
        )
    return ruchy


def okno(ruchy: list[Ruch]) -> tuple[datetime, datetime]:
    """Okres bierze się z pliku: pierwszy i ostatni ruch w eksporcie."""
    momenty = [r.moment for r in ruchy]
    return min(momenty), max(momenty)


# ---------------------------------------------------------------------------
# Zestawienie
# ---------------------------------------------------------------------------
def _lokalnie(moment: datetime) -> str:
    return moment.astimezone(REPORT_TZ).strftime("%Y-%m-%d %H:%M")


def _podsumuj(ruchy: list[Ruch]) -> dict:
    """Ile ruchów, kiedy pierwszy i ostatni, jakie operacje."""
    momenty = [r.moment for r in ruchy]
    kody: dict[str, int] = {}
    for r in ruchy:
        kody[r.code] = kody.get(r.code, 0) + 1
    return {
        "moves": len(ruchy),
        "first": _lokalnie(min(momenty)),
        "last": _lokalnie(max(momenty)),
        "codes": sorted(kody.items(), key=lambda p: (-p[1], p[0])),
    }


def zgrupuj(ruchy: list[Ruch]) -> dict[str, list[Ruch]]:
    """Ruchy po loginie. Klucz wielkimi literami — WMS i Pinokio nie zawsze
    piszą login tak samo, a to ten sam człowiek."""
    po_loginie: dict[str, list[Ruch]] = {}
    for r in ruchy:
        po_loginie.setdefault(r.login.upper(), []).append(r)
    return po_loginie


async def okresy_posiadania(db: AsyncSession, od: datetime, do: datetime) -> dict[int, list]:
    """Kiedy kto miał sprzęt w rękach — pracownik → lista przedziałów czasu.

    Zwrot nie niesie pracownika (karty nikt wtedy nie przykłada), więc stan
    odtwarzamy po urządzeniu: wydanie otwiera przedział, każda następna
    transakcja na tym samym sprzęcie go zamyka.

    Zaczynamy od stanu sprzed okresu, bo zmiana nocna pobiera skanery
    wieczorem — godziny przed pierwszym ruchem z pliku.
    """
    # Stan na początek okresu: ostatnia transakcja każdego urządzenia przed nim.
    wczesniej = (
        select(TransactionDB.device_id,
               func.max(TransactionDB.timestamp).label("kiedy"))
        .where(TransactionDB.timestamp < od)
        .group_by(TransactionDB.device_id)
        .subquery()
    )
    poczatkowe = (await db.execute(
        select(TransactionDB)
        .join(wczesniej, (TransactionDB.device_id == wczesniej.c.device_id)
              & (TransactionDB.timestamp == wczesniej.c.kiedy))
    )).scalars().all()

    wewnatrz = (await db.execute(
        select(TransactionDB)
        .where(TransactionDB.timestamp >= od, TransactionDB.timestamp <= do)
        .order_by(TransactionDB.timestamp)
    )).scalars().all()

    trzyma: dict[int, tuple[int, datetime]] = {}      # device_id → (pracownik, od kiedy)
    okresy: dict[int, list[tuple[datetime, datetime]]] = {}

    def zamknij(device_id: int, koniec: datetime) -> None:
        wpis = trzyma.pop(device_id, None)
        if wpis:
            okresy.setdefault(wpis[0], []).append((wpis[1], koniec))

    for transakcja in poczatkowe:
        if (transakcja.type == TransactionType.registered
                and transakcja.employee_id is not None):
            trzyma[transakcja.device_id] = (transakcja.employee_id,
                                            _aware(transakcja.timestamp))

    for transakcja in wewnatrz:
        moment = _aware(transakcja.timestamp)
        zamknij(transakcja.device_id, moment)
        if (transakcja.type == TransactionType.registered
                and transakcja.employee_id is not None):
            trzyma[transakcja.device_id] = (transakcja.employee_id, moment)

    for device_id in list(trzyma):
        zamknij(device_id, do)

    return okresy


def _aware(moment: datetime) -> datetime:
    """Znacznik ze strefą. Postgres oddaje je zonowane, sqlite w testach nie."""
    return moment if moment.tzinfo else moment.replace(tzinfo=timezone.utc)


def mial_sprzet(okresy: list[tuple[datetime, datetime]], moment: datetime) -> bool:
    return any(od <= moment <= do for od, do in okresy)


async def zbierz(db: AsyncSession, ruchy: list[Ruch]) -> dict:
    od, do = okno(ruchy)
    po_loginie = zgrupuj(ruchy)

    pracownicy = (await db.execute(
        select(EmployeeDB).options(selectinload(EmployeeDB.site))
    )).scalars().all()
    znani = {(p.wms_login or "").upper(): p for p in pracownicy if p.wms_login}

    # Kiedy kto miał sprzęt w rękach — z tym porównujemy każdy ruch z osobna.
    okresy = await okresy_posiadania(db, od, do)

    # Kiedy w ogóle brał sprzęt ostatni raz — bez tego wiersz "bez rejestracji"
    # nic nie mówi o tym, czy to nowa osoba, czy ktoś, kto trzyma skaner od wczoraj.
    ostatnie = dict((await db.execute(
        select(TransactionDB.employee_id, func.max(TransactionDB.timestamp))
        .where(TransactionDB.type == TransactionType.registered)
        .group_by(TransactionDB.employee_id)
    )).all())

    bez_rejestracji, nieznani, zarejestrowanych = [], [], 0

    for login, jego_ruchy in po_loginie.items():
        podsumowanie = _podsumuj(jego_ruchy)
        osoba = znani.get(login)

        if osoba is None:
            nieznani.append({"login": jego_ruchy[0].login, **podsumowanie})
            continue

        # Sprzęt pobrany przed początkiem pliku też się liczy: zmiana nocna
        # bierze skanery wieczorem, a eksport zaczyna się po północy.
        jego_okresy = okresy.get(osoba.id, [])
        bez_sprzetu = [r for r in jego_ruchy if not mial_sprzet(jego_okresy, r.moment)]

        if not bez_sprzetu:
            zarejestrowanych += 1
            continue

        kiedys = ostatnie.get(osoba.id)
        bez_rejestracji.append({
            "login": osoba.wms_login,
            "first_name": osoba.first_name,
            "last_name": osoba.last_name,
            "company": osoba.company,
            "site": osoba.site.name if osoba.site else None,
            "moves_without": len(bez_sprzetu),
            "first_without": _lokalnie(min(r.moment for r in bez_sprzetu)),
            "last_without": _lokalnie(max(r.moment for r in bez_sprzetu)),
            "last_registration": _lokalnie(_aware(kiedys)) if kiedys else None,
            **podsumowanie,
        })

    # Najwięcej pracy bez sprzętu na górze — od tego zaczyna się pytanie.
    bez_rejestracji.sort(key=lambda w: w["moves_without"], reverse=True)
    nieznani.sort(key=lambda w: w["moves"], reverse=True)

    return {
        "report_name": REPORT_NAME,
        "generated_at": datetime.now(tz=REPORT_TZ).strftime("%Y-%m-%d %H:%M"),
        # Okres to zakres samego pliku — nikt go nie wpisuje ręcznie.
        "from": _lokalnie(od),
        "to": _lokalnie(do),
        "hours": round((do - od).total_seconds() / 3600, 1),
        "totals": {
            "rows_total": len(ruchy),
            "logins_in_window": len(po_loginie),
            "with_device": zarejestrowanych,
            "not_registered": len(bez_rejestracji),
            "unknown": len(nieznani),
        },
        "not_registered": bez_rejestracji,
        "unknown": nieznani,
    }


# ---------------------------------------------------------------------------
# Plik
# ---------------------------------------------------------------------------
KOLUMNY_BEZ = ("Login WMS", "Imię", "Nazwisko", "Firma", "Site",
               "Ruchy bez sprzętu", "Ruchy razem", "Pierwszy bez sprzętu",
               "Ostatni bez sprzętu", "Operacje", "Ostatnia rejestracja")
KOLUMNY_NIEZNANI = ("Login z WMS", "Ruchy w WMS", "Pierwszy ruch", "Ostatni ruch",
                    "Operacje")


def _kody_tekst(kody) -> str:
    return ", ".join(f"{kod} ×{ile}" for kod, ile in kody)


def _naglowek(sheet, kolumny) -> None:
    sheet.append(list(kolumny))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1E293B")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    for index, kolumna in enumerate(kolumny, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(kolumna) + 4)
    sheet.freeze_panes = "A2"


def build_workbook(raport: dict) -> bytes:
    workbook = Workbook()

    arkusz = workbook.active
    arkusz.title = "Bez sprzętu"
    _naglowek(arkusz, KOLUMNY_BEZ)
    for w in raport["not_registered"]:
        arkusz.append([
            w["login"], w["first_name"], w["last_name"], w["company"], w["site"] or "—",
            w["moves_without"], w["moves"], w["first_without"], w["last_without"],
            _kody_tekst(w["codes"]), w["last_registration"] or "nigdy",
        ])

    drugi = workbook.create_sheet("Nieznane loginy")
    _naglowek(drugi, KOLUMNY_NIEZNANI)
    for w in raport["unknown"]:
        drugi.append([w["login"], w["moves"], w["first"], w["last"],
                      _kody_tekst(w["codes"])])

    # Okres i liczby — żeby plik odczytany za tydzień nadal mówił, czego dotyczy.
    opis = workbook.create_sheet("Zakres")
    opis.append(["Raport", raport["report_name"]])
    opis.append(["Okres z pliku",
                 f"{raport['from']} – {raport['to']} ({raport['hours']} h)"])
    opis.append(["Wygenerowano", raport["generated_at"]])
    for klucz, wartosc in raport["totals"].items():
        opis.append([klucz, wartosc])
    opis.column_dimensions["A"].width = 24
    opis.column_dimensions["B"].width = 44

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def report_file_name(generated_on) -> str:
    return f"Praca_WMS_bez_sprzetu_{generated_on.isoformat()}.xlsx"


# ---------------------------------------------------------------------------
# Endpointy
# ---------------------------------------------------------------------------
async def _wczytaj(plik: UploadFile) -> list[Ruch]:
    raw = await plik.read()
    if len(raw) > MAX_FILE_BYTES:
        raise HTTPException(400, "Plik jest za duży (limit 25 MB)")
    return parse_wms(raw)


@router.post("/wms")
async def wms_report(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_admin),
):
    """Analiza wgranego eksportu z WMS. Okres wynika z pliku."""
    return await zbierz(db, await _wczytaj(file))


@router.post("/wms.xlsx")
async def wms_report_xlsx(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user=Depends(require_admin),
):
    """Ten sam wynik w pliku. Plik źródłowy wraca z przeglądarki drugi raz —
    nic nie trzymamy na serwerze, więc nie ma czego czyścić ani pilnować."""
    raport = await zbierz(db, await _wczytaj(file))
    nazwa = report_file_name(datetime.now(REPORT_TZ).date())

    return Response(
        content=build_workbook(raport),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(nazwa)},
    )
